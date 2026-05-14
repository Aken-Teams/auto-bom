"""File upload and parsing routes."""
import shutil
import time
import logging
from datetime import datetime
from pathlib import Path
from fastapi import APIRouter, UploadFile, File, Depends, HTTPException
from sqlalchemy.orm import Session
from app.database import get_db
from app.config import UPLOAD_DIR
from app.models.tables import UploadRecord, CanTemplate
from app.services.excel_parser import parse_bom_base, parse_can_template

log = logging.getLogger(__name__)

router = APIRouter(prefix="/api/upload", tags=["upload"])


@router.post("/bom-base")
def upload_bom_base(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Upload BOM base file (WXBMR005)."""
    t0 = time.time()
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Only Excel files are supported")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    save_path = UPLOAD_DIR / f"bom_base_{ts}_{file.filename}"
    with open(save_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    log.info(f"[bom-base] save file: {time.time()-t0:.2f}s")

    t1 = time.time()
    items = parse_bom_base(str(save_path))
    log.info(f"[bom-base] parse ({len(items)} rows): {time.time()-t1:.2f}s")

    t2 = time.time()
    record = UploadRecord(
        filename=file.filename,
        file_path=str(save_path),
        file_type="bom_base",
        row_count=len(items),
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    log.info(f"[bom-base] db write: {time.time()-t2:.2f}s | total: {time.time()-t0:.2f}s")

    return {
        "id": record.id,
        "filename": record.filename,
        "row_count": len(items),
        "items": items,
    }


@router.post("/can-template")
def upload_can_template(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Upload file containing 罐头 templates (C-CMAX导入清单 with 罐头 sheet)."""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Only Excel files are supported")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    save_path = UPLOAD_DIR / f"can_template_{ts}_{file.filename}"
    with open(save_path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    templates = parse_can_template(str(save_path))

    # Upsert into database
    for t in templates:
        existing = db.query(CanTemplate).filter_by(waf_code=t["waf_code"]).first()
        if existing:
            for k, v in t.items():
                setattr(existing, k, v)
        else:
            db.add(CanTemplate(**t))

    # Save upload record
    record = UploadRecord(
        filename=file.filename,
        file_path=str(save_path),
        file_type="can_template",
        row_count=len(templates),
    )
    db.add(record)
    db.commit()
    db.refresh(record)

    return {"id": record.id, "count": len(templates), "templates": templates}


@router.post("/std-operations")
def upload_std_operations(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Upload WXBMR004 standard operations file.

    Saves file and does lightweight validation only.
    Full parsing happens on-demand during file generation.
    """
    t0 = time.time()
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Only Excel files are supported")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    save_path = UPLOAD_DIR / f"std_ops_{ts}_{file.filename}"
    with open(save_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    log.info(f"[std-ops] save file: {time.time()-t0:.2f}s")

    # Lightweight validation: just open the file to confirm it's a valid Excel
    t1 = time.time()
    try:
        from openpyxl import load_workbook as _lw
        _wb = _lw(str(save_path), read_only=True)
        total_rows = max((_wb.active.max_row or 1) - 1, 0)
        _wb.close()
    except Exception:
        save_path.unlink(missing_ok=True)
        raise HTTPException(400, "Invalid file. Please upload the correct WXBMR004 Excel file.")
    log.info(f"[std-ops] validate ({total_rows} rows): {time.time()-t1:.2f}s")

    t2 = time.time()
    record = UploadRecord(
        filename=file.filename,
        file_path=str(save_path),
        file_type="std_operation",
        row_count=total_rows,
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    log.info(f"[std-ops] db write: {time.time()-t2:.2f}s | total: {time.time()-t0:.2f}s")

    return {"id": record.id, "filename": record.filename, "row_count": total_rows}


@router.get("/records")
def list_upload_records(db: Session = Depends(get_db)):
    """List all upload records."""
    records = db.query(UploadRecord).order_by(UploadRecord.uploaded_at.desc()).all()
    return [
        {
            "id": r.id,
            "filename": r.filename,
            "file_type": r.file_type,
            "row_count": r.row_count,
            "uploaded_at": r.uploaded_at.isoformat() if r.uploaded_at else None,
        }
        for r in records
    ]


@router.get("/can-templates")
def list_can_templates(db: Session = Depends(get_db)):
    """List all can templates in database."""
    templates = db.query(CanTemplate).all()
    return [
        {
            "id": t.id,
            "function": t.function,
            "waf_code": t.waf_code,
            "supplier": t.supplier,
            "wafer_size": t.wafer_size,
            "mil": t.mil,
            "weld_can": t.weld_can,
            "weld_desc": t.weld_desc,
            "mold_can": t.mold_can,
            "mold_desc": t.mold_desc,
            "pack_can": t.pack_can,
            "pack_desc": t.pack_desc,
        }
        for t in templates
    ]
